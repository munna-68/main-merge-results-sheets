#!/usr/bin/env python3
"""One-time fixture authoring.

Converts the two MIS 3201 Word mark files into .xlsx fixtures (Excel-only input is a
product decision) and builds the student-master workbook from the published transcript.

The .xlsx outputs are committed, so this only needs re-running if the fixtures are
rebuilt from scratch. The Word sources were deleted from the working tree once the
fixtures were verified against the department's own result sheet; restore them first:

    git show 1da977b:'MIS 3201 Continous.docx' > 'MIS 3201 Continous.docx'
    git show 1da977b:'MIS 3201 Final.docx' > 'MIS 3201 Final.docx'
    git show 1da977b:'3rd Year 2nd Semester Final Result-2025 (1).docx' \
        > '3rd Year 2nd Semester Final Result-2025 (1).docx'

Then:  python3 scripts/author-fixtures.py
"""
import os
import re
import sys

import docx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "test", "fixtures", "inputs")

BENGALI_DIGITS = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")


def require(name):
    path = os.path.join(ROOT, name)
    if not os.path.exists(path):
        raise SystemExit(
            f"missing source: {name}\n"
            "Restore it from git history first -- see the module docstring."
        )
    return path


def num(value):
    """Coerce a docx cell to a float, or None when it is not a number."""
    if value is None:
        return None
    text = str(value).translate(BENGALI_DIGITS).strip()
    text = text.replace("\u00a0", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def table_rows(doc):
    """Every table in the document, as a list of row-cell-lists."""
    for table in doc.tables:
        yield [[cell.text for cell in row.cells] for row in table.rows]


def find_student_table(doc):
    """(rows, id_col) for the marks table: the one densest in 8-digit ids."""
    best = None
    for rows in table_rows(doc):
        if not rows:
            continue
        width = max(len(r) for r in rows)
        if width < 3:
            continue
        for col in range(width):
            ids = 0
            for row in rows:
                if col < len(row) and re.fullmatch(
                    r"\d{6,}", row[col].translate(BENGALI_DIGITS).strip()
                ):
                    ids += 1
            if ids >= 3 and (best is None or ids > best[2]):
                best = (rows, col, ids)
    if best is None:
        raise SystemExit("no marks table found")
    return best[0], best[1]


def header_index(rows, id_col):
    """The header is the first row whose id cell is not a number."""
    for i, row in enumerate(rows):
        cell = row[id_col] if id_col < len(row) else ""
        if not re.fullmatch(r"\d{6,}", cell.translate(BENGALI_DIGITS).strip()):
            return i
    return 0


def write_book(path, header, rows, widths):
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10, name="Times New Roman")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"wrote {os.path.relpath(path, ROOT)}  ({len(rows)} rows)")


def convert_marks(source_name, target_name):
    doc = docx.Document(require(source_name))
    rows, id_col = find_student_table(doc)
    h = header_index(rows, id_col)
    header = [c.strip() for c in rows[h]]
    body = rows[h + 1:]

    parsed = []
    for row in body:
        sid = num(row[id_col]) if id_col < len(row) else None
        if sid is None or sid < 100000:
            continue
        marks = [num(row[i]) for i in range(id_col + 1, len(row))]
        while marks and marks[-1] is None:
            marks.pop()
        parsed.append([int(sid)] + marks)

    width = max(len(r) for r in parsed)
    # Keep the document's own column labels. Flattening them to Q1..Qn would
    # hide the fact that this is a continuous-assessment sheet, which is how the
    # app tells one apart from a setter paper.
    source_labels = header[id_col + 1:]
    while source_labels and num(source_labels[-1]) is None and not re.search(
        r"total|মোট", source_labels[-1], re.I
    ):
        source_labels.pop()
    header_out = ["ID"] + [str(s).strip() for s in source_labels[: width - 2]] + ["Total"]

    for row in parsed:
        while len(row) < len(header_out):
            row.append(None)

    write_book(
        os.path.join(OUT, target_name),
        header_out,
        parsed,
        [12] + [7] * (len(header_out) - 1),
    )
    return parsed, header


IDENTITY = re.compile(
    r"Registration\s*No\s*[:.]?\s*(\d+)"
    r".*?Session\s*[:.]?\s*([0-9]{4}\s*-\s*[0-9]{4})"
    r".*?Name\s*[:.]?\s*(.+?)"
    r"\s*Student\s*ID\s*[:.]?\s*(\d+)",
    re.IGNORECASE | re.DOTALL,
)


def author_master():
    """53 identities from the published 3rd-year transcript.

    Each student's identity sits in one paragraph of the form
    `Registration No : 000018177 Session : 2022-2023 Name : ... Student ID : 12222002`.
    """
    doc = docx.Document(require("3rd Year 2nd Semester Final Result-2025 (1).docx"))
    students = {}
    for para in doc.paragraphs:
        m = IDENTITY.search(" ".join(para.text.split()))
        if not m:
            continue
        reg, session, name, sid = m.groups()
        students[int(sid)] = {
            "id": str(int(sid)),
            "name": name.strip(),
            "reg": str(int(reg)).zfill(9),
            "session": session.replace(" ", ""),
        }

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["SL", "ID", "Name", "Reg", "Session"])
    for i, sid in enumerate(sorted(students), start=1):
        s = students[sid]
        ws.append([i, s["id"], s["name"], s["reg"], s["session"]])
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10, name="Times New Roman")
    for width, col in [(6, "A"), (12, "B"), (32, "C"), (14, "D"), (12, "E")]:
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    path = os.path.join(OUT, "student-master.xlsx")
    wb.save(path)
    print(f"wrote test/fixtures/inputs/student-master.xlsx  ({len(students)} students)")
    return students


if __name__ == "__main__":
    cont, ch = convert_marks("MIS 3201 Continous.docx", "MIS 3201 Continuous.xlsx")
    fin, fh = convert_marks("MIS 3201 Final.docx", "MIS 3201 Final.xlsx")
    print(f"  continuous header: {ch}")
    print(f"  final header:      {fh}")
    author_master()